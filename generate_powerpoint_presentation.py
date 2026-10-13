import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Color Palette Definitions
    DARK_NAVY = RGBColor(10, 25, 49)      # #0A1931 - Primary Brand Navy
    BRAND_RED = RGBColor(217, 4, 41)      # #D90429 - Accent Red
    BRAND_GOLD = RGBColor(245, 158, 11)   # #F59E0B - Accent Gold/Amber
    EMERALD = RGBColor(16, 185, 129)      # #10B981 - Success Green
    SLATE_DARK = RGBColor(30, 41, 59)     # #1E293B - Dark Text
    SLATE_MUTED = RGBColor(100, 116, 139) # #64748B - Muted Text
    LIGHT_BG = RGBColor(248, 249, 250)    # #F8F9FA - Slide Background
    CARD_BG = RGBColor(255, 255, 255)     # #FFFFFF - Card Background
    CARD_BORDER = RGBColor(226, 232, 240) # #E2E8F0 - Card Outline
    BLUE_ACCENT = RGBColor(37, 99, 235)   # #2563EB - Tech Blue

    def set_slide_background(slide, color):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = color
        bg.line.fill.background()
        return bg

    def add_header(slide, category_tag, title_text, is_dark=False):
        # Category Tag Badge
        tag_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(8.0), Inches(0.35))
        tf_tag = tag_box.text_frame
        tf_tag.word_wrap = True
        tf_tag.margin_left = tf_tag.margin_top = tf_tag.margin_right = tf_tag.margin_bottom = 0
        p_tag = tf_tag.paragraphs[0]
        p_tag.text = category_tag.upper()
        p_tag.font.name = "Segoe UI"
        p_tag.font.size = Pt(11)
        p_tag.font.bold = True
        p_tag.font.color.rgb = BRAND_GOLD if is_dark else BRAND_RED

        # Main Slide Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.75), Inches(11.5), Inches(0.7))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        tf_title.margin_left = tf_title.margin_top = tf_title.margin_right = tf_title.margin_bottom = 0
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.name = "Segoe UI"
        p_title.font.size = Pt(24)
        p_title.font.bold = True
        p_title.font.color.rgb = CARD_BG if is_dark else DARK_NAVY

    def add_card(slide, left, top, width, height, bg_color=CARD_BG, border_color=CARD_BORDER, border_width=1):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = bg_color
        if border_color:
            shape.line.color.rgb = border_color
            shape.line.width = Pt(border_width)
        else:
            shape.line.fill.background()
        return shape

    def add_bullet_list(text_frame, items, font_size=13, line_spacing=1.2, text_color=SLATE_DARK):
        for i, item in enumerate(items):
            p = text_frame.add_paragraph() if i > 0 or len(text_frame.paragraphs[0].text) > 0 else text_frame.paragraphs[0]
            p.text = item
            p.font.name = "Segoe UI"
            p.font.size = Pt(font_size)
            p.font.color.rgb = text_color
            p.space_after = Pt(8)

    # ==========================================================
    # SLIDE 1: TITLE SLIDE (DARK NAVY)
    # ==========================================================
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1, DARK_NAVY)

    # Decorative red accent bar
    bar = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.2), Inches(0.12), Inches(5.0))
    bar.fill.solid()
    bar.fill.fore_color.rgb = BRAND_RED
    bar.line.fill.background()

    # Logo if available
    logo_path = "public/images/logo_white.png"
    if os.path.exists(logo_path):
        s1.shapes.add_picture(logo_path, Inches(1.2), Inches(1.2), height=Inches(1.2))

    # Title & Subtitle Box
    t_box = s1.shapes.add_textbox(Inches(1.2), Inches(2.6), Inches(11.0), Inches(2.5))
    tf1 = t_box.text_frame
    tf1.word_wrap = True
    
    p = tf1.paragraphs[0]
    p.text = "JUBIS MARKETING"
    p.font.name = "Segoe UI"
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = CARD_BG
    
    p2 = tf1.add_paragraph()
    p2.text = "Enterprise B2B ERP & Client E-Commerce Portal"
    p2.font.name = "Segoe UI"
    p2.font.size = Pt(24)
    p2.font.color.rgb = BRAND_GOLD
    p2.space_before = Pt(8)

    p3 = tf1.add_paragraph()
    p3.text = "Comprehensive System Feature Architecture & Operational Walkthrough"
    p3.font.name = "Segoe UI"
    p3.font.size = Pt(15)
    p3.font.color.rgb = RGBColor(203, 213, 225)
    p3.space_before = Pt(12)

    # Live Cloud Badge Card
    badge_card = add_card(s1, Inches(1.2), Inches(5.3), Inches(10.8), Inches(1.0), bg_color=RGBColor(15, 34, 64), border_color=BRAND_GOLD)
    b_box = s1.shapes.add_textbox(Inches(1.4), Inches(5.4), Inches(10.4), Inches(0.8))
    b_tf = b_box.text_frame
    b_p1 = b_tf.paragraphs[0]
    b_p1.text = "LIVE CLOUD PRODUCTION URL:"
    b_p1.font.name = "Segoe UI"
    b_p1.font.size = Pt(11)
    b_p1.font.bold = True
    b_p1.font.color.rgb = BRAND_GOLD
    
    b_p2 = b_tf.add_paragraph()
    b_p2.text = "https://jubis-marketing-production-m2g1zm.laravel.cloud/"
    b_p2.font.name = "Segoe UI"
    b_p2.font.size = Pt(15)
    b_p2.font.bold = True
    b_p2.font.color.rgb = CARD_BG

    # ==========================================================
    # SLIDE 2: EXECUTIVE SUMMARY & BUSINESS BACKGROUND
    # ==========================================================
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_background(s2, LIGHT_BG)
    add_header(s2, "Executive Overview", "Bridging B2B Electrical Distribution & Digital ERP Automation")

    # Card 1: The B2B Problem Statement
    add_card(s2, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2))
    c1_box = s2.shapes.add_textbox(Inches(1.1), Inches(1.8), Inches(5.0), Inches(4.7))
    tf = c1_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Traditional Industry Bottlenecks"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = BRAND_RED
    p.space_after = Pt(14)
    add_bullet_list(tf, [
        "Fragmented Quotations: B2B clients rely on manual phone calls, Viber chats, and paper RFQs for fluctuating electrical prices.",
        "Phantom Stock & Inventory Inaccuracy: High-SKU warehouses frequently suffer from overselling and unrecorded adjustments.",
        "Delayed Invoicing & Credit Reconciliation: Offline billing leads to friction in collecting receivables and tracking Net-30 credit.",
        "Zero Logistics Traceability: Customers have no real-time visibility into whether wholesale deliveries are dispatched or in transit."
    ], font_size=13)

    # Card 2: The Jubis ERP Solution
    add_card(s2, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2))
    c2_box = s2.shapes.add_textbox(Inches(7.1), Inches(1.8), Inches(5.1), Inches(4.7))
    tf = c2_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "The Integrated Jubis ERP Solution"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = DARK_NAVY
    p.space_after = Pt(14)
    add_bullet_list(tf, [
        "Self-Service Client RFQ Portal: B2B buyers browse dynamic electrical catalogs, configure technical variants, and submit quotes directly.",
        "Role-Based Departmental Automation: Dedicated workspaces for Sales, Finance, Purchasing, Warehouse, and Executive Super Admins.",
        "Real-Time Immutable Inventory Audit: Live stock reservation upon quote approval and automatic PO stock increments.",
        "Dual Settlement Engine: Digital checkout via PayMongo (GCash, Maya, Cards) and automated Net-30 credit terms for bulk orders.",
        "Live Delivery & Warranty Tracking: Built-in delivery tracking simulator and structured RMA (Return Merchandise Authorization) workflows."
    ], font_size=13)

    # ==========================================================
    # SLIDE 3: SYSTEM ARCHITECTURE & FULL TECH STACK
    # ==========================================================
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_background(s3, LIGHT_BG)
    add_header(s3, "System Architecture", "Modern Multi-Tier Enterprise Cloud Architecture")

    # Add Diagram Image
    diagram_tech = "public/images/presentation/diagram_tech_stack.png"
    if os.path.exists(diagram_tech):
        s3.shapes.add_picture(diagram_tech, Inches(0.8), Inches(1.5), width=Inches(11.73))

    # ==========================================================
    # SLIDE 4: DEPARTMENTAL HIERARCHY & ROLE-BASED ACCESS CONTROL
    # ==========================================================
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_background(s4, LIGHT_BG)
    add_header(s4, "Governance & Security", "Departmental Role-Based Access Control (RBAC)")

    diagram_rbac = "public/images/presentation/diagram_rbac.png"
    if os.path.exists(diagram_rbac):
        s4.shapes.add_picture(diagram_rbac, Inches(0.8), Inches(1.5), width=Inches(11.73))

    # ==========================================================
    # SLIDE 5: END-TO-END ENTERPRISE WORKFLOW (RFQ TO FULFILLMENT)
    # ==========================================================
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_background(s5, LIGHT_BG)
    add_header(s5, "Operational Workflow", "Complete Enterprise Lifecycle: From Catalog to Delivery")

    diagram_wf = "public/images/presentation/diagram_workflow.png"
    if os.path.exists(diagram_wf):
        s5.shapes.add_picture(diagram_wf, Inches(0.8), Inches(1.5), width=Inches(11.73))

    # ==========================================================
    # SLIDE 6: CLIENT SIDE — PUBLIC PORTAL & BRAND CATALOG
    # ==========================================================
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_background(s6, LIGHT_BG)
    add_header(s6, "Client Portal", "Product Discovery, Smart Search & Faceted Filtering")

    # Left Column: Features Card
    add_card(s6, Inches(0.8), Inches(1.6), Inches(7.0), Inches(5.2))
    box = s6.shapes.add_textbox(Inches(1.1), Inches(1.8), Inches(6.4), Inches(4.7))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "E-Commerce Catalog Capabilities"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = DARK_NAVY
    p.space_after = Pt(14)
    add_bullet_list(tf, [
        "Real-Time Keyword Search: Instant search across thousands of SKUs, brand names, and technical product descriptions with URL query synchronization.",
        "Multi-Faceted Category Filter: Filter through Wire & Cables, Control Components, Circuit Protection, Lighting, Industrial Connectors, and Hardware.",
        "Brand Priority Sorting: Intelligent catalog ranking prioritizing strategic partners (OMNI, Jubis, Schneider Electric, Phelps Dodge, Fuji, Philips).",
        "Live Inventory Status Indicators: Real-time badge display ('In Stock', 'Out of Stock') preventing client frustration before quote generation.",
        "Corporate Storytelling & Inquiries: Dedicated About Us CMS and Contact Inquiry form with throttled mail notifications to Jubis management."
    ], font_size=13)

    # Right Column: Product Visuals
    add_card(s6, Inches(8.1), Inches(1.6), Inches(4.4), Inches(2.45))
    img_thhn = "public/images/products/THHNWire.png"
    if os.path.exists(img_thhn):
        s6.shapes.add_picture(img_thhn, Inches(8.3), Inches(1.75), height=Inches(1.8))
    cap_box = s6.shapes.add_textbox(Inches(10.2), Inches(1.9), Inches(2.1), Inches(1.5))
    cap_tf = cap_box.text_frame
    cap_tf.word_wrap = True
    p = cap_tf.paragraphs[0]
    p.text = "Phelps Dodge THHN"
    p.font.bold = True
    p.font.size = Pt(13)
    p = cap_tf.add_paragraph()
    p.text = "Building Wire\nHigh-spec industrial cable with nylon jacket."
    p.font.size = Pt(11)
    p.font.color.rgb = SLATE_MUTED

    add_card(s6, Inches(8.1), Inches(4.35), Inches(4.4), Inches(2.45))
    img_led = "public/images/products/LEDbulb.png"
    if os.path.exists(img_led):
        s6.shapes.add_picture(img_led, Inches(8.3), Inches(4.5), height=Inches(1.8))
    cap_box2 = s6.shapes.add_textbox(Inches(10.2), Inches(4.65), Inches(2.1), Inches(1.5))
    cap_tf2 = cap_box2.text_frame
    cap_tf2.word_wrap = True
    p = cap_tf2.paragraphs[0]
    p.text = "Philips LED Fixture"
    p.font.bold = True
    p.font.size = Pt(13)
    p = cap_tf2.add_paragraph()
    p.text = "Commercial Lighting\nEnergy-efficient industrial bulb & ceiling luminaires."
    p.font.size = Pt(11)
    p.font.color.rgb = SLATE_MUTED

    # ==========================================================
    # SLIDE 7: CLIENT SIDE — TECHNICAL SPECS & VARIANT SELECTOR
    # ==========================================================
    s7 = prs.slides.add_slide(blank_layout)
    set_slide_background(s7, LIGHT_BG)
    add_header(s7, "Client Portal", "Engineering Specifications & Dynamic Variant Selector")

    add_card(s7, Inches(0.8), Inches(1.6), Inches(7.0), Inches(5.2))
    box = s7.shapes.add_textbox(Inches(1.1), Inches(1.8), Inches(6.4), Inches(4.7))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Engineering-Grade Product Detail Pages"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = DARK_NAVY
    p.space_after = Pt(14)
    add_bullet_list(tf, [
        "Comprehensive Spec Sheets: Dedicated parameter tables for electrical engineers (Voltage, Amperage, Wire Gauge, Phase, Certifications).",
        "Dynamic Variant Switching: Switch between package types, coil voltages, pole counts, and sizes without full page reloads.",
        "Tiered Wholesale Pricing: Transparent unit wholesale rates with dynamic volume calculations for large-scale engineering contractors.",
        "Stock Availability Verification: Checks specific SKU variant warehouse counts to ensure fulfillment readiness.",
        "Related Product Carousel: Recommends matching components (e.g. matching circuit breaker enclosures, lugs, or contactor auxiliary relays)."
    ], font_size=13)

    add_card(s7, Inches(8.1), Inches(1.6), Inches(4.4), Inches(2.45))
    img_mc = "public/images/products/magneticcontractor.png"
    if os.path.exists(img_mc):
        s7.shapes.add_picture(img_mc, Inches(8.3), Inches(1.75), height=Inches(1.8))
    cap_box = s7.shapes.add_textbox(Inches(10.2), Inches(1.9), Inches(2.1), Inches(1.5))
    cap_tf = cap_box.text_frame
    cap_tf.word_wrap = True
    p = cap_tf.paragraphs[0]
    p.text = "Fuji Magnetic Contactor"
    p.font.bold = True
    p.font.size = Pt(13)
    p = cap_tf.add_paragraph()
    p.text = "Industrial Control\n32A 220V heavy-duty motor starter switch."
    p.font.size = Pt(11)
    p.font.color.rgb = SLATE_MUTED

    add_card(s7, Inches(8.1), Inches(4.35), Inches(4.4), Inches(2.45))
    img_cb = "public/images/products/circuitbreaker.png"
    if os.path.exists(img_cb):
        s7.shapes.add_picture(img_cb, Inches(8.3), Inches(4.5), height=Inches(1.8))
    cap_box2 = s7.shapes.add_textbox(Inches(10.2), Inches(4.65), Inches(2.1), Inches(1.5))
    cap_tf2 = cap_box2.text_frame
    cap_tf2.word_wrap = True
    p = cap_tf2.paragraphs[0]
    p.text = "Schneider Breaker"
    p.font.bold = True
    p.font.size = Pt(13)
    p = cap_tf2.add_paragraph()
    p.text = "Circuit Protection\n63A thermal-magnetic overload protection."
    p.font.size = Pt(11)
    p.font.color.rgb = SLATE_MUTED

    # ==========================================================
    # SLIDE 8: CLIENT SIDE — INTERACTIVE RFQ CART & BULK QUOTING
    # ==========================================================
    s8 = prs.slides.add_slide(blank_layout)
    set_slide_background(s8, LIGHT_BG)
    add_header(s8, "Client Portal", "Interactive RFQ Cart & Wholesale Quote Submission")

    add_card(s8, Inches(0.8), Inches(1.6), Inches(7.0), Inches(5.2))
    box = s8.shapes.add_textbox(Inches(1.1), Inches(1.8), Inches(6.4), Inches(4.7))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "B2B Request-For-Quote (RFQ) Model"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = DARK_NAVY
    p.space_after = Pt(14)
    add_bullet_list(tf, [
        "B2B Quoting vs Consumer Checkout: Replaces rigid retail checkout with an RFQ cart designed for commercial negotiation and volume pricing.",
        "Interactive Line Item Quantity Controls: Immediate subtotal recalculation, VAT estimation, and quantity validation against warehouse availability.",
        "Custom Project Notes: Contractors can attach site delivery requirements, project reference codes, and billing instructions.",
        "Inactive SKU Shield: Strict API guards ensure unlisted or discontinued products cannot be injected into the quotation cart.",
        "Automatic Pipeline Routing: Submitting the cart automatically converts it from 'Draft' to 'Pending' in the Sales Department workbench."
    ], font_size=13)

    add_card(s8, Inches(8.1), Inches(1.6), Inches(4.4), Inches(2.45))
    img_sock = "public/images/products/Industrial Socket.png"
    if os.path.exists(img_sock):
        s8.shapes.add_picture(img_sock, Inches(8.3), Inches(1.75), height=Inches(1.8))
    cap_box = s8.shapes.add_textbox(Inches(10.2), Inches(1.9), Inches(2.1), Inches(1.5))
    cap_tf = cap_box.text_frame
    cap_tf.word_wrap = True
    p = cap_tf.paragraphs[0]
    p.text = "PCE Industrial Socket"
    p.font.bold = True
    p.font.size = Pt(13)
    p = cap_tf.add_paragraph()
    p.text = "Heavy-Duty Connector\n3-Pin 32A weatherproof socket assembly."
    p.font.size = Pt(11)
    p.font.color.rgb = SLATE_MUTED

    add_card(s8, Inches(8.1), Inches(4.35), Inches(4.4), Inches(2.45))
    img_fuse = "public/images/products/Glass Fuse.png"
    if os.path.exists(img_fuse):
        s8.shapes.add_picture(img_fuse, Inches(8.3), Inches(4.5), height=Inches(1.8))
    cap_box2 = s8.shapes.add_textbox(Inches(10.2), Inches(4.65), Inches(2.1), Inches(1.5))
    cap_tf2 = cap_box2.text_frame
    cap_tf2.word_wrap = True
    p = cap_tf2.paragraphs[0]
    p.text = "Bussman Glass Fuse"
    p.font.bold = True
    p.font.size = Pt(13)
    p = cap_tf2.add_paragraph()
    p.text = "High-Voltage Protection\nFast-acting 10A 250V overcurrent fuse."
    p.font.size = Pt(11)
    p.font.color.rgb = SLATE_MUTED

    # ==========================================================
    # SLIDE 9: CLIENT SIDE — CLIENT DASHBOARD & ORDER LIFECYCLE
    # ==========================================================
    s9 = prs.slides.add_slide(blank_layout)
    set_slide_background(s9, LIGHT_BG)
    add_header(s9, "Client Portal", "Client Dashboard, Quote Statuses & Order Lifecycle")

    # Card 1: Dashboard & Lifecycle
    add_card(s9, Inches(0.8), Inches(1.6), Inches(5.7), Inches(5.2))
    box = s9.shapes.add_textbox(Inches(1.1), Inches(1.8), Inches(5.1), Inches(4.7))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Customer Dashboard Experience"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = DARK_NAVY
    p.space_after = Pt(14)
    add_bullet_list(tf, [
        "Authenticated User Hub (/dashboard): Clients access active quotes, invoices, and shipment tracking in one consolidated screen.",
        "Quote Lifecycle Status Badges:\n• Pending: Submitted, awaiting sales pricing review.\n• Reviewed: Sales has appraised and added custom discounts.\n• Approved: Officially finalized, unlocking invoice & payment.\n• Rejected: Notification with inventory release.",
        "Transparent Financial Overview: Shows total quoted values, VAT breakdowns, payment status, and due dates.",
        "Order Historical Records: Permanent repository of past orders for repetitive corporate purchasing and procurement audits."
    ], font_size=13)

    # Card 2: Interactive Actions
    add_card(s9, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2))
    box2 = s9.shapes.add_textbox(Inches(7.1), Inches(1.8), Inches(5.1), Inches(4.7))
    tf2 = box2.text_frame
    tf2.word_wrap = True
    p = tf2.paragraphs[0]
    p.text = "Direct Client Action Triggers"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = BRAND_RED
    p.space_after = Pt(14)
    add_bullet_list(tf2, [
        "Instant Invoice PDF Download: One-click export of official stamped Jubis Marketing PDF invoices for procurement documentation.",
        "Online Payment Gateway: Click 'Pay Now' to launch the secure PayMongo digital payment interface.",
        "Net-30 Credit Terms Activation: High-volume wholesale orders unlock Net-30 credit terms at the touch of a button.",
        "Live Carrier Tracking: Track logistics delivery progress in real-time with dispatched carrier updates.",
        "RMA Warranty Returns: Easily submit defective component return requests directly against specific invoice lines."
    ], font_size=13)

    # ==========================================================
    # SLIDE 10: CLIENT SIDE — INVOICING, PAYMONGO & NET-30 CREDIT
    # ==========================================================
    s10 = prs.slides.add_slide(blank_layout)
    set_slide_background(s10, LIGHT_BG)
    add_header(s10, "Client Portal", "Dual Settlement Engine: PayMongo Gateway & Net-30 Terms")

    # Card 1: PayMongo Gateway
    add_card(s10, Inches(0.8), Inches(1.6), Inches(5.7), Inches(5.2))
    box = s10.shapes.add_textbox(Inches(1.1), Inches(1.8), Inches(5.1), Inches(4.7))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "PayMongo Digital Payments"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = BLUE_ACCENT
    p.space_after = Pt(14)
    add_bullet_list(tf, [
        "Omnichannel Payment Options: Direct support for GCash, Maya, QR Ph, and Mastercard/Visa credit and debit cards.",
        "Centavo-Accurate Float Calculations: Eliminates binary floating-point rounding errors when submitting payment totals.",
        "Automated Payment Link Dispatch: Sales quote approval automatically generates a custom PayMongo checkout link.",
        "Replay-Resistant Webhook Architecture: Payment callbacks are protected by HMAC-SHA256 signature and timestamp verification.",
        "Sandbox Payment Simulator: Built-in mock payment environment enabling real-time demonstration without live financial risk."
    ], font_size=13)

    # Card 2: Enterprise Net-30 Credit
    add_card(s10, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2))
    box2 = s10.shapes.add_textbox(Inches(7.1), Inches(1.8), Inches(5.1), Inches(4.7))
    tf2 = box2.text_frame
    tf2.word_wrap = True
    p = tf2.paragraphs[0]
    p.text = "Enterprise Net-30 Credit Terms"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = BRAND_GOLD
    p.space_after = Pt(14)
    add_bullet_list(tf2, [
        "Volume-Based Eligibility Guard: Available exclusively to large wholesale purchases exceeding 1,000 units.",
        "30-Day Automated Terms: Sets official invoice status to 'on_terms' and schedules automated 30-day receivable maturity.",
        "Immediate Order Release: Clears invoices for immediate warehouse picking and shipping prior to cash settlement.",
        "Anti-Tamper Status Guard: Prevents already paid invoices from being improperly reverted into credit terms.",
        "AR Ledger Synchronization: Instantly records outstanding receivables in the Finance Department reporting dashboard."
    ], font_size=13)

    # ==========================================================
    # SLIDE 11: CLIENT SIDE — SHIPMENT TRACKING & RMA RETURNS
    # ==========================================================
    s11 = prs.slides.add_slide(blank_layout)
    set_slide_background(s11, LIGHT_BG)
    add_header(s11, "Client Portal", "Real-Time Logistics Tracking & Warranty RMA Processing")

    add_card(s11, Inches(0.8), Inches(1.6), Inches(7.0), Inches(5.2))
    box = s11.shapes.add_textbox(Inches(1.1), Inches(1.8), Inches(6.4), Inches(4.7))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Logistics Visibility & Warranty Protection"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = DARK_NAVY
    p.space_after = Pt(14)
    add_bullet_list(tf, [
        "Live Delivery Tracking Simulator (/api/mock/tracking/{id}): Full tracking page displaying dispatch time, assigned carrier, driver details, and current parcel location.",
        "Lalamove Delivery Integration: Supports third-party courier APIs as well as in-house Jubis delivery fleet dispatch.",
        "RMA Warranty Claims Workflow: Clients can initiate return merchandise requests directly against verified paid invoice lines.",
        "Itemized Defect Categorization: Select from 'Defective Unit', 'Damaged in Transit', or 'Wrong Specification' with customer notes.",
        "Atomic Database Rollback: Strict database transactions guarantee that partial line failures cleanly roll back parent RMA records without orphaned data."
    ], font_size=13)

    add_card(s11, Inches(8.1), Inches(1.6), Inches(4.4), Inches(2.45))
    img_fan = "public/images/products/Ceiling or Wall Exhaust Fan.png"
    if os.path.exists(img_fan):
        s11.shapes.add_picture(img_fan, Inches(8.3), Inches(1.75), height=Inches(1.8))
    cap_box = s11.shapes.add_textbox(Inches(10.2), Inches(1.9), Inches(2.1), Inches(1.5))
    cap_tf = cap_box.text_frame
    cap_tf.word_wrap = True
    p = cap_tf.paragraphs[0]
    p.text = "Exhaust Ventilation"
    p.font.bold = True
    p.font.size = Pt(13)
    p = cap_tf.add_paragraph()
    p.text = "Ceiling & Wall Fan\nCommercial ventilation fan with warranty protection."
    p.font.size = Pt(11)
    p.font.color.rgb = SLATE_MUTED

    add_card(s11, Inches(8.1), Inches(4.35), Inches(4.4), Inches(2.45))
    img_wp = "public/images/products/outdoor electrical weatherproof cover.png"
    if os.path.exists(img_wp):
        s11.shapes.add_picture(img_wp, Inches(8.3), Inches(4.5), height=Inches(1.8))
    cap_box2 = s11.shapes.add_textbox(Inches(10.2), Inches(4.65), Inches(2.1), Inches(1.5))
    cap_tf2 = cap_box2.text_frame
    cap_tf2.word_wrap = True
    p = cap_tf2.paragraphs[0]
    p.text = "Weatherproof Cover"
    p.font.bold = True
    p.font.size = Pt(13)
    p = cap_tf2.add_paragraph()
    p.text = "Outdoor Electrical Box\nHeavy-duty IP65 weatherproof protective cover."
    p.font.size = Pt(11)
    p.font.color.rgb = SLATE_MUTED

    # ==========================================================
    # SLIDE 12: ADMIN ERP — EXECUTIVE DASHBOARD & ANALYTICS
    # ==========================================================
    s12 = prs.slides.add_slide(blank_layout)
    set_slide_background(s12, DARK_NAVY)
    add_header(s12, "Admin ERP Command Center", "Executive KPI Dashboard & Real-Time Financial Analytics", is_dark=True)

    # 4 Metric Cards
    metric_w = Inches(2.75)
    metric_h = Inches(1.5)
    m_y = Inches(1.6)
    m_gap = Inches(0.24)

    metrics_data = [
        ("TOTAL REVENUE", "PHP 2,450,800", "Paid Invoice Settlements", EMERALD),
        ("OUTSTANDING AR", "PHP 640,250", "Unpaid & On-Terms Balance", BRAND_GOLD),
        ("PENDING QUOTES", "14 Active RFQs", "Awaiting Sales Review", BLUE_ACCENT),
        ("LOW STOCK SKUs", "8 Critical Items", "Below Reorder Threshold", BRAND_RED)
    ]

    for i, (m_title, m_val, m_sub, m_col) in enumerate(metrics_data):
        x = Inches(0.8) + i * (metric_w + m_gap)
        add_card(s12, x, m_y, metric_w, metric_h, bg_color=RGBColor(15, 34, 64), border_color=m_col, border_width=1.5)
        m_box = s12.shapes.add_textbox(x + Inches(0.2), m_y + Inches(0.15), metric_w - Inches(0.4), metric_h - Inches(0.3))
        mtf = m_box.text_frame
        mtf.word_wrap = True
        p = mtf.paragraphs[0]
        p.text = m_title
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = m_col
        
        p = mtf.add_paragraph()
        p.text = m_val
        p.font.size = Pt(18)
        p.font.bold = True
        p.font.color.rgb = CARD_BG
        p.space_before = Pt(4)

        p = mtf.add_paragraph()
        p.text = m_sub
        p.font.size = Pt(9)
        p.font.color.rgb = RGBColor(148, 163, 184)
        p.space_before = Pt(2)

    # Lower Split Card
    add_card(s12, Inches(0.8), Inches(3.35), Inches(5.75), Inches(3.5), bg_color=RGBColor(15, 34, 64), border_color=CARD_BORDER)
    b_box1 = s12.shapes.add_textbox(Inches(1.1), Inches(3.55), Inches(5.2), Inches(3.1))
    tf1 = b_box1.text_frame
    tf1.word_wrap = True
    p = tf1.paragraphs[0]
    p.text = "6-Month Revenue Analytics (Chart.js)"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = BRAND_GOLD
    p.space_after = Pt(10)
    add_bullet_list(tf1, [
        "Dynamic Monthly Revenue Engine: Aggregates settled revenue across rolling 6-month cycles directly from the invoice ledger.",
        "Cash Flow Visibility: Differentiates instant digital settlements from Net-30 receivables collection.",
        "Interactive Chart.js Dashboard: Responsive chart component with hover tooltips and financial breakdown."
    ], font_size=12, text_color=RGBColor(226, 232, 240))

    add_card(s12, Inches(6.8), Inches(3.35), Inches(5.75), Inches(3.5), bg_color=RGBColor(15, 34, 64), border_color=CARD_BORDER)
    b_box2 = s12.shapes.add_textbox(Inches(7.1), Inches(3.55), Inches(5.2), Inches(3.1))
    tf2 = b_box2.text_frame
    tf2.word_wrap = True
    p = tf2.paragraphs[0]
    p.text = "Real-Time Actionable Feeds"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = BRAND_RED
    p.space_after = Pt(10)
    add_bullet_list(tf2, [
        "Recent Quotations Stream: Immediate visibility into arriving RFQs with client company name and total line items.",
        "Recent Invoices Feed: Real-time settlement monitoring for accounting staff with instant payment status pills.",
        "Urgent Reorder Alert Table: Prioritized list of inventory items with quantities < 20 units for instant PO generation."
    ], font_size=12, text_color=RGBColor(226, 232, 240))

    # ==========================================================
    # SLIDE 13: ADMIN ERP — INVENTORY CONTROL & AUDIT TRAIL
    # ==========================================================
    s13 = prs.slides.add_slide(blank_layout)
    set_slide_background(s13, LIGHT_BG)
    add_header(s13, "Admin ERP Inventory", "Inventory Control, Stock Movements & CSV Export")

    add_card(s13, Inches(0.8), Inches(1.6), Inches(7.0), Inches(5.2))
    box = s13.shapes.add_textbox(Inches(1.1), Inches(1.8), Inches(6.4), Inches(4.7))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Warehouse Inventory Management Workbench"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = DARK_NAVY
    p.space_after = Pt(14)
    add_bullet_list(tf, [
        "Master Product Catalog Table: Complete oversight of 900+ SKUs across all electrical categories with pagination and search.",
        "Active State Toggle: Instantly enable or disable products from appearing in the client-facing catalog.",
        "Immutable Stock Movement Ledger (stock_movements): Every single inventory delta logs:\n• Exact quantity change (+ or -)\n• Actor User ID and staff role\n• Movement type ('purchase', 'adjustment', 'order')\n• Reference tag (e.g. PO-12, QUOTE-REJECT-5).",
        "Safe CSV Inventory Export: Sanitized one-click CSV export shielding against spreadsheet formula injection attacks (prefixes '=', '+', '-', '@').",
        "Zero Phantom Stock Guarantee: Reverting or rejecting quotes automatically restores warehouse stock with atomic logging."
    ], font_size=13)

    add_card(s13, Inches(8.1), Inches(1.6), Inches(4.4), Inches(5.2))
    img_lugs = "public/images/products/generated_lugs.jpg"
    if os.path.exists(img_lugs):
        s13.shapes.add_picture(img_lugs, Inches(8.3), Inches(1.85), height=Inches(2.4))
    side_box = s13.shapes.add_textbox(Inches(8.3), Inches(4.5), Inches(4.0), Inches(2.1))
    stf = side_box.text_frame
    stf.word_wrap = True
    p = stf.paragraphs[0]
    p.text = "Granular SKU Control"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = BRAND_RED
    p.space_after = Pt(6)
    p = stf.add_paragraph()
    p.text = "Supports heavy industrial connectors, cables, transformers, and copper cable lugs with real-time warehouse count tracking and reorder limits."
    p.font.size = Pt(12)
    p.font.color.rgb = SLATE_MUTED

    # ==========================================================
    # SLIDE 14: ADMIN ERP — PROCUREMENT & SUPPLIER MANAGEMENT (PO)
    # ==========================================================
    s14 = prs.slides.add_slide(blank_layout)
    set_slide_background(s14, LIGHT_BG)
    add_header(s14, "Admin ERP Procurement", "Supplier Directory & Purchase Order (PO) Engine")

    add_card(s14, Inches(0.8), Inches(1.6), Inches(5.7), Inches(5.2))
    box = s14.shapes.add_textbox(Inches(1.1), Inches(1.8), Inches(5.1), Inches(4.7))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Supplier Directory Management"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = DARK_NAVY
    p.space_after = Pt(14)
    add_bullet_list(tf, [
        "Centralized Vendor Profiles (/admin/suppliers): Comprehensive database of electrical manufacturers and component importers.",
        "Supplier Compliance Metadata: Records official Tax Identification Number (TIN), corporate addresses, and primary sales contacts.",
        "Direct Procurement Communications: Phone numbers and verified billing emails stored for automated replenishment requests.",
        "Supplier Purchase Order History: Instant visibility into all past and active purchase orders linked to each vendor profile."
    ], font_size=13)

    add_card(s14, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2))
    box2 = s14.shapes.add_textbox(Inches(7.1), Inches(1.8), Inches(5.1), Inches(4.7))
    tf2 = box2.text_frame
    tf2.word_wrap = True
    p = tf2.paragraphs[0]
    p.text = "Purchase Order (PO) Lifecycle"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = EMERALD
    p.space_after = Pt(14)
    add_bullet_list(tf2, [
        "Multi-Line PO Workbench: Create purchase orders with multiple product SKUs, supplier unit costs, and delivery date commitments.",
        "Lifecycle State Machine: Transitions cleanly through Draft ➔ Ordered ➔ Partially Received ➔ Received.",
        "Automatic Stock Replenishment: Transitioning PO to 'Received' automatically increments physical warehouse stock and logs movement.",
        "Negative Stock Prevention: If a received PO is reverted back to ordered, the system checks whether inventory is sufficient to prevent negative stock.",
        "Real-Time Cost of Goods Tracking: Calculates exact landed cost per unit for accurate margin calculation."
    ], font_size=13)

    # ==========================================================
    # SLIDE 15: ADMIN ERP — SALES, FINANCE & LOGISTICS WORKBENCHES
    # ==========================================================
    s15 = prs.slides.add_slide(blank_layout)
    set_slide_background(s15, LIGHT_BG)
    add_header(s15, "Admin ERP Operations", "Sales Quotes, Financial Ledger & Shipment Dispatch")

    col_w = Inches(3.75)
    col_h = Inches(5.2)
    gap = Inches(0.24)

    # Col 1: Sales
    add_card(s15, Inches(0.8), Inches(1.6), col_w, col_h)
    box1 = s15.shapes.add_textbox(Inches(1.0), Inches(1.8), col_w - Inches(0.4), col_h - Inches(0.4))
    tf1 = box1.text_frame
    tf1.word_wrap = True
    p = tf1.paragraphs[0]
    p.text = "Sales & Quotes"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = BRAND_RED
    p.space_after = Pt(10)
    add_bullet_list(tf1, [
        "Quote Review Queue: Inspect client RFQs with customer corporate details.",
        "Wholesale Price Adjustments: Apply contractor discounts or custom shipping fees.",
        "One-Click Approval: Automatically spawns official invoice, sets Net-30 date, and generates PayMongo link.",
        "Stock Reservation Protection: Approving reserves stock; rejecting cleanly restores warehouse quantities."
    ], font_size=12)

    # Col 2: Finance
    add_card(s15, Inches(0.8) + col_w + gap, Inches(1.6), col_w, col_h)
    box2 = s15.shapes.add_textbox(Inches(1.0) + col_w + gap, Inches(1.8), col_w - Inches(0.4), col_h - Inches(0.4))
    tf2 = box2.text_frame
    tf2.word_wrap = True
    p = tf2.paragraphs[0]
    p.text = "Finance & Invoicing"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = BRAND_GOLD
    p.space_after = Pt(10)
    add_bullet_list(tf2, [
        "Comprehensive Invoice Ledger: Filter by unpaid, partially paid, paid, and overdue.",
        "Manual Payment Logging: Post bank wire or cash checks with date and reference.",
        "Overpayment Defense Guard: Automatically rejects payment attempts exceeding the outstanding balance.",
        "VAT & Tax Reporting (/admin/reports): Official breakdown of 12% output VAT, taxable sales, and exempt transactions."
    ], font_size=12)

    # Col 3: Logistics
    add_card(s15, Inches(0.8) + (col_w + gap) * 2, Inches(1.6), col_w, col_h)
    box3 = s15.shapes.add_textbox(Inches(1.0) + (col_w + gap) * 2, Inches(1.8), col_w - Inches(0.4), col_h - Inches(0.4))
    tf3 = box3.text_frame
    tf3.word_wrap = True
    p = tf3.paragraphs[0]
    p.text = "Logistics & Shipments"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = EMERALD
    p.space_after = Pt(10)
    add_bullet_list(tf3, [
        "Order Fulfillment Queue: Shows all paid or credit-approved invoices ready for dispatch.",
        "Carrier Allocation: Assign shipping carriers (Lalamove, In-House Jubis Trucks, 3PL).",
        "Tracking Number Generation: Attaches unique tracking IDs accessible by clients.",
        "Automated Delivery Mails: Changing shipment status to 'delivered' triggers automated delivery confirmation emails to client."
    ], font_size=12)

    # ==========================================================
    # SLIDE 16: ADMIN ERP — SYSTEM SETTINGS & QUALITY GATES
    # ==========================================================
    s16 = prs.slides.add_slide(blank_layout)
    set_slide_background(s16, LIGHT_BG)
    add_header(s16, "Admin Governance & QA", "System Configuration, Security Controls & Release Gates")

    add_card(s16, Inches(0.8), Inches(1.6), Inches(5.7), Inches(5.2))
    box = s16.shapes.add_textbox(Inches(1.1), Inches(1.8), Inches(5.1), Inches(4.7))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Enterprise System Administration"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = DARK_NAVY
    p.space_after = Pt(14)
    add_bullet_list(tf, [
        "Company Profile Settings: Manage company legal name, official address, TIN, and contact details reflected on invoices.",
        "Configurable VAT Engine: Dynamic VAT percentage configuration (default 12%) with decimal precision validation.",
        "Strict Field Whitelisting: Rejects unknown secret fields (e.g. SMTP passwords) from ever being cached or exposed publicly.",
        "Staff User Administration: Super Admins can add, edit roles, and assign departments across all employees."
    ], font_size=13)

    add_card(s16, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2))
    box2 = s16.shapes.add_textbox(Inches(7.1), Inches(1.8), Inches(5.1), Inches(4.7))
    tf2 = box2.text_frame
    tf2.word_wrap = True
    p = tf2.paragraphs[0]
    p.text = "Enterprise Security & Quality Gates"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = EMERALD
    p.space_after = Pt(14)
    add_bullet_list(tf2, [
        "100% Passed Release Gates: Comprehensive PHPUnit test suite validating all P0 & P1 operational boundary cases.",
        "Database Constraint Safeguards: Enforces strict unique constraints preventing duplicate invoices or shipments.",
        "Replay-Resistant Webhook Verification: PayMongo webhooks validate HMAC signatures and reject stale requests (>1 hour).",
        "Role-Based Route Isolation: Granular middleware prevents privilege escalation between departments."
    ], font_size=13)

    # ==========================================================
    # SLIDE 17: CONCLUSION & LIVE DEMONSTRATION HANDOVER
    # ==========================================================
    s17 = prs.slides.add_slide(blank_layout)
    set_slide_background(s17, DARK_NAVY)
    add_header(s17, "Project Demonstration", "Live Cloud Handover & System Verification", is_dark=True)

    add_card(s17, Inches(0.8), Inches(1.6), Inches(5.7), Inches(5.2), bg_color=RGBColor(15, 34, 64), border_color=BRAND_GOLD)
    box1 = s17.shapes.add_textbox(Inches(1.1), Inches(1.8), Inches(5.1), Inches(4.7))
    tf1 = box1.text_frame
    tf1.word_wrap = True
    p = tf1.paragraphs[0]
    p.text = "Cloud Production Hosting"
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = BRAND_GOLD
    p.space_after = Pt(14)
    add_bullet_list(tf1, [
        "Production URL:\nhttps://jubis-marketing-production-m2g1zm.laravel.cloud/",
        "Infrastructure Provider: Laravel Cloud",
        "Region: Asia Pacific (Singapore) Edge Network",
        "Database Engine: MySQL 8.4 Enterprise Cluster",
        "Features: SSL Automated Encryption, Edge CDN, Scale-to-Zero Compute & Database Architecture.",
        "Zero Downtime: Multi-session support allows multiple concurrent users to test across departments simultaneously."
    ], font_size=13, text_color=CARD_BG)

    add_card(s17, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2), bg_color=RGBColor(15, 34, 64), border_color=EMERALD)
    box2 = s17.shapes.add_textbox(Inches(7.1), Inches(1.8), Inches(5.1), Inches(4.7))
    tf2 = box2.text_frame
    tf2.word_wrap = True
    p = tf2.paragraphs[0]
    p.text = "Live Demonstration Agenda"
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = EMERALD
    p.space_after = Pt(14)
    add_bullet_list(tf2, [
        "1. Client Flow: Browse catalog, filter by OMNI / Phelps Dodge, and submit an RFQ cart.",
        "2. Sales Review: Log in as Sales / CEO, review RFQ lines, adjust pricing, and approve quote.",
        "3. Financial Settlement: Inspect generated invoice, trigger PayMongo simulation or Net-30 credit terms.",
        "4. Warehouse & Dispatch: Verify stock deduction and create a shipment with carrier tracking.",
        "5. Client Tracking & RMA: View live shipment tracking page and submit a warranty RMA return request."
    ], font_size=13, text_color=CARD_BG)

    output_filename = "Jubis_Marketing_ERP_Features_Presentation.pptx"
    prs.save(output_filename)
    print(f"Successfully generated {output_filename} with {len(prs.slides)} slides!")

if __name__ == "__main__":
    create_deck()
